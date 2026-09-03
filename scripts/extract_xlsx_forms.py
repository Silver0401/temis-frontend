#!/usr/bin/env python3
"""Extract reviewable dynamic-form metadata from one or more XLSX templates.

Usage:
  python3 -m pip install openpyxl
  python3 scripts/extract_xlsx_forms.py
  python3 scripts/extract_xlsx_forms.py file1.xlsx file2.xlsx --output result.json

The default inputs are the two Temis workbooks in ~/Downloads/Temis Data.
Source files are only read and are never copied or modified.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
import warnings
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries


DEFAULT_INPUTS = [
    Path.home() / "Downloads/Temis Data/FORMATOS EXPEDIENTE CLINICO 2025.xlsx",
    Path.home() / "Downloads/Temis Data/Formatos Aviso MP_2025.xlsx",
]
DEFAULT_OUTPUT = Path(__file__).with_name("temis-form-inputs.json")

SECTION_RE = re.compile(r"^(?:[IVXLCDM]+\.\s+|[A-Z]\)\s*)(.+)", re.IGNORECASE)
INLINE_CHECK_RE = re.compile(r"([A-ZÁÉÍÓÚÜÑ][A-ZÁÉÍÓÚÜÑ .ÁÉÍÓÚÜ/()-]{1,45}):\s*\(\s*\)", re.IGNORECASE)
NUMBERED_OPTION_RE = re.compile(r"^\d+[.)]\s*(.+)$")
LABEL_RE = re.compile(r"([^:\n]{2,80}):")

DATE_WORDS = ("FECHA", "FUM", "FUP", "FUA", "FUC", "PRÓXIMA CITA", "PROXIMA CITA")
NUMBER_WORDS = (
    "EDAD",
    "PESO",
    "TALLA",
    "IMC",
    "PERÍMETRO",
    "PERIMETRO",
    "FOLIO",
    "CAMA",
    "DOSIS",
    "INGRESO",
    "NÚMERO",
    "NUMERO",
    "NO. ",
    "FREC.",
    "TEMPERATURA",
    "GLUCOSA",
    "GLICEMIA",
)
TEXTAREA_WORDS = (
    "MOTIVO",
    "PADECIMIENTO",
    "EVOLUCIÓN",
    "EVOLUCION",
    "EXPLORACIÓN",
    "EXPLORACION",
    "IMPRESIÓN DIAGNÓSTICA",
    "IMPRESION DIAGNOSTICA",
    "PRONÓSTICO",
    "PRONOSTICO",
    "CUIDADOS GENERALES",
    "ESTUDIOS AUXILIARES",
    "INTERVENCIONES",
    "RESULTADO DE LA VISITA",
    "PLAN",
)
STATIC_PREFIXES = (
    "SERVICIOS DE SALUD",
    "SECRETARÍA DE SALUD",
    "SECRETARIA DE SALUD",
    "JURISDICCIÓN SANITARIA",
    "JURISDICCION SANITARIA",
)


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized(value: str) -> str:
    return "".join(
        char
        for char in unicodedata.normalize("NFD", clean(value).upper())
        if unicodedata.category(char) != "Mn"
    )


def slug(value: str) -> str:
    result = re.sub(r"[^a-z0-9]+", "_", normalized(value).lower()).strip("_")
    return result[:64] or "field"


def unique_identifier(base: str, used: Counter[str]) -> str:
    used[base] += 1
    return base if used[base] == 1 else f"{base}_{used[base]}"


def printed_options(value: Any) -> list[str]:
    raw = str(value).strip()
    if len(raw) > 80 or any(character in raw for character in ".,("):
        return [clean(raw)]
    return [clean(part) for part in re.split(r"\s{2,}", raw) if clean(part)]


def input_type(label: str, has_large_region: bool = False) -> tuple[str, str, str]:
    upper = normalized(label)
    if any(word in upper for word in DATE_WORDS):
        return "date", "medium", "La etiqueta impresa solicita una fecha"
    if any(word in upper for word in NUMBER_WORDS):
        return "number", "medium", "La etiqueta o unidad impresa solicita una cantidad"
    if has_large_region or any(word in upper for word in TEXTAREA_WORDS):
        return "textarea", "medium", "El formato reserva un área narrativa"
    return "text", "medium", "El formato reserva una línea o celda libre"


def merged_range_for(sheet: Any, row: int, column: int) -> str | None:
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return str(merged)
    return None


def has_fillable_neighbor(sheet: Any, row: int, column: int) -> tuple[bool, bool, list[str]]:
    coordinates: list[str] = []
    large = False
    for col in range(column + 1, min(sheet.max_column, column + 4) + 1):
        cell = sheet.cell(row, col)
        if cell.value is None and cell.has_style and not isinstance(cell, MergedCell):
            coordinates.append(cell.coordinate)
            merged = merged_range_for(sheet, row, col)
            if merged:
                min_col, min_row, max_col, max_row = range_boundaries(merged)
                large = large or max_col - min_col >= 2 or max_row - min_row >= 1
            break
    for next_row in range(row + 1, min(sheet.max_row, row + 3) + 1):
        cell = sheet.cell(next_row, column)
        if cell.value is None and cell.has_style and not isinstance(cell, MergedCell):
            coordinates.append(cell.coordinate)
            merged = merged_range_for(sheet, next_row, column)
            if merged:
                min_col, min_row, max_col, max_row = range_boundaries(merged)
                large = large or max_col - min_col >= 2 or max_row - min_row >= 1
            break
    return bool(coordinates), large, coordinates


def section_title(text: str) -> str | None:
    match = SECTION_RE.match(clean(text))
    return clean(match.group(1)) if match else None


def nearest_label(sheet: Any, row: int, column: int) -> str:
    for col in range(column - 1, 0, -1):
        value = clean(sheet.cell(row, col).value)
        if value:
            return value.rstrip(":_")
    for previous_row in range(row - 1, max(0, row - 4), -1):
        value = clean(sheet.cell(previous_row, column).value)
        if value:
            return value.rstrip(":_")
    return f"Campo {sheet.cell(row, column).coordinate}"


def validation_options(workbook: Any, sheet: Any, formula: str | None) -> list[str]:
    if not formula:
        return []
    formula = formula.lstrip("=")
    if formula.startswith('"') and formula.endswith('"'):
        return [clean(option) for option in formula[1:-1].split(",") if clean(option)]

    target_sheet = sheet
    target_range = formula
    if "!" in formula:
        sheet_name, target_range = formula.rsplit("!", 1)
        target_sheet = workbook[sheet_name.strip("'")]
    elif formula in workbook.defined_names:
        destinations = list(workbook.defined_names[formula].destinations)
        if not destinations:
            return []
        sheet_name, target_range = destinations[0]
        target_sheet = workbook[sheet_name]

    try:
        cells = target_sheet[target_range.replace("$", "")]
    except (KeyError, ValueError):
        return []
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    elif cells and not isinstance(cells[0], tuple):
        cells = (cells,)
    return [clean(cell.value) for row in cells for cell in row if clean(cell.value)]


def source(sheet: Any, cells: Iterable[str], evidence: str, confidence: str) -> dict[str, Any]:
    return {
        "sheet": sheet.title,
        "cells": list(cells),
        "evidence": evidence,
        "confidence": confidence,
    }


def make_input(
    identifier: str,
    label: str,
    kind: str,
    section: str,
    source_data: dict[str, Any],
    *,
    required: bool | None = None,
    options: list[str] | None = None,
) -> dict[str, Any]:
    return {
        "identifier": identifier,
        "label": clean(label).strip(":_ "),
        "type": kind,
        "required": required,
        "options": options or [],
        "section": section,
        "source": source_data,
    }


def extract_validations(workbook: Any, sheet: Any, used: Counter[str]) -> tuple[list[dict[str, Any]], set[str]]:
    inputs: list[dict[str, Any]] = []
    covered: set[str] = set()
    validations = list(sheet.data_validations.dataValidation) if sheet.data_validations else []
    for validation in validations:
        options = validation_options(workbook, sheet, validation.formula1)
        kind = "select" if validation.type == "list" else validation.type
        if kind not in {"select", "date", "decimal", "whole", "textLength", "custom"}:
            kind = "text"
        if kind in {"decimal", "whole"}:
            kind = "number"
        if kind in {"textLength", "custom"}:
            kind = "text"
        for cell_range in validation.sqref.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(cell_range))
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    cell = sheet.cell(row, column)
                    label = nearest_label(sheet, row, column)
                    identifier = unique_identifier(slug(label), used)
                    inputs.append(
                        make_input(
                            identifier,
                            label,
                            kind,
                            "Validaciones de Excel",
                            source(sheet, [cell.coordinate], f"dataValidation {validation.formula1 or ''}".strip(), "high"),
                            required=not validation.allow_blank,
                            options=options,
                        )
                    )
                    covered.add(cell.coordinate)
    return inputs, covered


OPTION_LABELS = (
    "TIPO DE UNIDAD",
    "ESTADO NUTRICIONAL",
    "CLASIFICACION",
    "PLAN TERAPEUTICO",
    "ESTADO GENERAL",
    "OJOS",
    "BOCA Y LENGUA",
    "RESPIRACION",
    "SED",
    "ELASTICIDAD DE LA PIEL",
    "PULSO",
    "LLENADO CAPILAR",
    "FONTANELA",
    "RIESGO DE CAIDA",
)


def option_heading_above(sheet: Any, row: int, column: int) -> tuple[str, int] | None:
    for previous_row in range(row - 1, max(0, row - 4), -1):
        for previous_col in range(sheet.max_column, 0, -1):
            text = clean(sheet.cell(previous_row, previous_col).value)
            if any(
                word in normalized(text)
                for word in ("CLASIFICACION", "PLAN TERAPEUTICO", "SELECCIONE", "ESCALA")
            ):
                return text.split(":", 1)[0].split("(", 1)[0].strip(": "), previous_col
    return None


def explicit_row_group(sheet: Any, row: int) -> tuple[str, list[str], list[str], str] | None:
    cells = [cell for cell in sheet[row] if clean(cell.value)]
    if len(cells) < 3:
        return None
    for index, cell in enumerate(cells):
        label = clean(cell.value).rstrip(":")
        if not any(word in normalized(label) for word in OPTION_LABELS):
            continue
        option_cells = cells[index + 1 :]
        options = [
            part
            for option_cell in option_cells
            for part in printed_options(option_cell.value)
            if clean(part) and ":" not in clean(part)
        ]
        if options and normalized(options[0]) in {"KG", "CM", "KG/M2", "MMHG", "LPM", "RPM", "PPM"}:
            continue
        if len(options) >= 2:
            return label, [option.strip("_ ") for option in options], [cell.coordinate, *[item.coordinate for item in option_cells]], "radio"

    first = clean(cells[0].value).rstrip(":")
    heading = option_heading_above(sheet, row, cells[0].column)
    if not heading:
        return None
    heading_label, heading_column = heading
    option_cells = [cell for cell in cells if cell.column >= heading_column]
    options = [
        part
        for cell in option_cells
        for part in printed_options(cell.value)
        if clean(part) and ":" not in clean(part) and normalized(part) != "SOMATOMETRIA"
    ]
    if len(options) < 2:
        return None
    kind = "checkbox" if normalized(heading_label).startswith("SELECCIONE") else "radio"
    return heading_label, [option.strip("_ ") for option in options], [cell.coordinate for cell in option_cells], kind


def extract_numbered_groups(sheet: Any, used: Counter[str]) -> tuple[list[dict[str, Any]], set[str]]:
    inputs: list[dict[str, Any]] = []
    consumed: set[str] = set()
    for column in range(1, sheet.max_column + 1):
        candidates = []
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            match = NUMBERED_OPTION_RE.match(clean(cell.value))
            if match and "\n" not in str(cell.value):
                candidates.append((row, cell, clean(match.group(1))))

        group: list[tuple[int, Any, str]] = []
        for candidate in [*candidates, (10**9, None, "")]:
            if group and candidate[0] - group[-1][0] > 2:
                if len(group) >= 2:
                    first_row = group[0][0]
                    nearby = []
                    for previous_row in range(first_row - 1, max(0, first_row - 4), -1):
                        for previous_col in range(1, sheet.max_column + 1):
                            text = clean(sheet.cell(previous_row, previous_col).value)
                            if text.endswith(":"):
                                nearby.append((abs(column - previous_col), text.rstrip(":")))
                    label = min(nearby, default=(0, f"Selección {group[0][1].coordinate}"))[1]
                    cells = [item[1].coordinate for item in group]
                    inputs.append(
                        make_input(
                            unique_identifier(slug(label), used),
                            label,
                            "radio",
                            "General",
                            source(sheet, cells, "Opciones numeradas consecutivas", "high"),
                            options=[item[2] for item in group],
                        )
                    )
                    consumed.update(cells)
                group = []
            group.append(candidate)
    return inputs, consumed


def extract_sheet(workbook: Any, sheet: Any) -> dict[str, Any]:
    used: Counter[str] = Counter()
    inputs, validation_cells = extract_validations(workbook, sheet, used)
    numbered_inputs, numbered_cells = extract_numbered_groups(sheet, used)
    inputs.extend(numbered_inputs)
    consumed: set[str] = set(validation_cells) | numbered_cells
    current_section = "General"

    for row_index in range(1, sheet.max_row + 1):
        values = [(cell, clean(cell.value)) for cell in sheet[row_index] if clean(cell.value)]
        if not values:
            continue

        for _, text in values:
            section = section_title(text)
            if section:
                current_section = section
                break

        group = explicit_row_group(sheet, row_index)
        if group:
            label, options, coordinates, kind = group
            identifier = unique_identifier(slug(label), used)
            inputs.append(
                make_input(
                    identifier,
                    label,
                    kind,
                    current_section,
                    source(sheet, coordinates, "Opciones enumeradas en la misma fila", "high"),
                    options=options,
                )
            )
            consumed.update(coordinates)

        inline_checks: list[tuple[Any, str]] = []
        for cell, text in values:
            for match in INLINE_CHECK_RE.finditer(text):
                inline_checks.append((cell, clean(match.group(1))))
        if len(inline_checks) >= 2:
            parent = clean(values[0][1]).split(":", 1)[0]
            options = [label for _, label in inline_checks]
            coordinates = sorted({cell.coordinate for cell, _ in inline_checks})
            inputs.append(
                make_input(
                    unique_identifier(slug(parent), used),
                    parent,
                    "checkbox",
                    current_section,
                    source(sheet, coordinates, "Opciones con casillas impresas '( )'", "high"),
                    options=options,
                )
            )
            consumed.update(coordinates)

        for cell, text in values:
            if cell.coordinate in consumed or normalized(text).startswith(STATIC_PREFIXES):
                continue

            section = section_title(text)
            bare_section = clean(section or "")
            has_neighbor, large_region, neighbors = has_fillable_neighbor(
                sheet, cell.row, cell.column
            )

            inline_yes_no = re.search(r"\bSI\b\s*_+\s*\bNO\b", normalized(text))
            if inline_yes_no:
                label = text.split(":", 1)[0]
                inputs.append(
                    make_input(
                        unique_identifier(slug(label), used),
                        label,
                        "radio",
                        current_section,
                        source(sheet, [cell.coordinate], "Alternativas SI/NO impresas", "high"),
                        options=["SI", "NO"],
                    )
                )
                consumed.add(cell.coordinate)
                continue

            inline_labels = [clean(match.group(1)).strip("_ ") for match in LABEL_RE.finditer(text)]
            if len(text) < 140 and len(inline_labels) >= 3:
                inputs.append(
                    make_input(
                        unique_identifier(slug(inline_labels[0]), used),
                        inline_labels[0],
                        "radio",
                        current_section,
                        source(sheet, [cell.coordinate], "Alternativas impresas dentro de la misma celda", "high"),
                        options=inline_labels[1:],
                    )
                )
                consumed.add(cell.coordinate)
                continue

            label_matches = [clean(match.group(1)) for match in LABEL_RE.finditer(text)]
            labels = label_matches if label_matches else []
            if not labels and "_" in text:
                labels = [clean(text.split("_", 1)[0]).rstrip(":")]
            if not labels and section and any(word in normalized(section) for word in TEXTAREA_WORDS):
                labels = [bare_section]
                large_region = True
            if len(text) > 140 and not section:
                continue
            if not labels or (not has_neighbor and ":" not in text and "_" not in text and not large_region):
                continue

            for label in labels:
                if len(label) < 2 or normalized(label).startswith(STATIC_PREFIXES):
                    continue
                kind, confidence, evidence = input_type(label, large_region)
                inputs.append(
                    make_input(
                        unique_identifier(slug(label), used),
                        label,
                        kind,
                        current_section,
                        source(sheet, [cell.coordinate, *neighbors], evidence, confidence),
                    )
                )
            consumed.add(cell.coordinate)

    by_section: dict[str, list[dict[str, Any]]] = {}
    for item in inputs:
        by_section.setdefault(item["section"], []).append(item)
    sections = [{"name": name, "inputs": items} for name, items in by_section.items()]
    counts = Counter(item["type"] for item in inputs)
    validation_count = len(sheet.data_validations.dataValidation) if sheet.data_validations else 0
    return {
        "name": sheet.title,
        "sheet_state": sheet.sheet_state,
        "dimensions": sheet.calculate_dimension(),
        "data_validation_count": validation_count,
        "field_count": len(inputs),
        "field_types": dict(sorted(counts.items())),
        "sections": sections,
    }


def extract_workbook(path: Path) -> dict[str, Any]:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support is incomplete")
        workbook = load_workbook(path, data_only=False, read_only=False)
    forms = [extract_sheet(workbook, sheet) for sheet in workbook.worksheets]
    totals = Counter()
    for form in forms:
        totals.update(form["field_types"])
    return {
        "source_file": str(path),
        "sheet_count": len(forms),
        "field_count": sum(form["field_count"] for form in forms),
        "field_types": dict(sorted(totals.items())),
        "forms": forms,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="*", type=Path, default=DEFAULT_INPUTS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    missing = [str(path) for path in args.inputs if not path.exists()]
    if missing:
        parser.error(f"No existen: {', '.join(missing)}")

    result = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "purpose": "Revisión previa de inputs para formularios dinámicos; no integrado a Temis",
        "supported_types": ["text", "textarea", "select", "date", "number", "checkbox", "radio"],
        "limitations": [
            "Estos libros no contienen data validations, nombres definidos ni controles nativos.",
            "required=null significa que el archivo no especifica obligatoriedad.",
            "Los tipos semánticos se respaldan con texto, unidades o geometría impresa y deben revisarse antes de integrar.",
            "DrawingML e imágenes no se convierten automáticamente en controles.",
        ],
        "workbooks": [extract_workbook(path.resolve()) for path in args.inputs],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"JSON: {args.output}")
    for workbook in result["workbooks"]:
        summary = ", ".join(f"{kind}={count}" for kind, count in workbook["field_types"].items())
        print(f"{Path(workbook['source_file']).name}: {workbook['field_count']} campos ({summary})")


if __name__ == "__main__":
    main()
